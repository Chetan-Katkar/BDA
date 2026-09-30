import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

K = json.load(open('verified_final.json'))
R = json.load(open('removed.json'))
FK = ['1 HR/recruitment email','2 General company email','3 LinkedIn page','4 Careers page','5 Current opening','6 Prior intern/fresher evidence']
HDR = PatternFill('solid', fgColor='1F3864'); HF = Font(bold=True, color='FFFFFF', size=9)
THIN = Border(*[Side(style='thin', color='D0D0D0')]*4)
YES = PatternFill('solid', fgColor='C6EFCE'); NO = PatternFill('solid', fgColor='F2F2F2')
wb = Workbook()

def style(ws, hrow, ncol, widths):
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
        c = ws.cell(hrow, ci); c.fill = HDR; c.font = HF
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[hrow].height = 40
    for row in ws.iter_rows(min_row=hrow+1, max_row=ws.max_row, max_col=ncol):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True); c.border = THIN; c.font = Font(size=9)

# Sheet 1: the six-point checklist (the answer to the question asked)
ws = wb.active; ws.title = 'Final 75 - Six Point Check'
ws['A1'] = ('FINAL VERIFIED LIST - 75 companies. For each one, the six columns say what you actually have in hand. '
            'Sorted so companies with a verified HR/recruitment email come first, then by how many of the six you have. '
            'Nobody has been contacted.')
ws.merge_cells('A1:M1'); ws['A1'].font = Font(bold=True, size=11)
ws['A1'].alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[1].height = 34
cols = ['No.','Company','Nearest Station','Area / Address'] + FK + ['Score /6','Emailable now','Company Type','Technology / Domain']
ws.append(cols)
for i, r in enumerate(K, 1):
    ws.append([i, r['company'], r['station'], r['area']] + [r['_f'][k] for k in FK]
              + [r['_ready'], r['_contactable'], r['type'], r['domain']])
style(ws, 2, len(cols), [5,34,17,44,13,13,11,11,11,14,9,11,24,38])
for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=5, max_col=10):
    for c in row:
        c.fill = YES if c.value == 'YES' else NO
        c.alignment = Alignment(horizontal='center', vertical='center')
        if c.value == 'YES': c.font = Font(size=9, bold=True)
ws.freeze_panes = 'B3'; ws.auto_filter.ref = f'A2:{get_column_letter(len(cols))}{ws.max_row}'

# Sheet 2: full detail
ws2 = wb.create_sheet('Final 75 - Full Detail')
d = [('No.',None,5),('Company','company',32),('Area / Address','area',44),('Nearest Station','station',16),
     ('Approx. Distance','dist',22),('Company Type','type',24),('Technology / Domain','domain',38),
     ('Relevant Internship Roles','roles',38),('Current Opening?','opening',30),('Internship Evidence','evidence',44),
     ('HR / Recruitment Email','hr_email',28),('General Email','gen_email',28),('Phone','phone',20),
     ('Website','website',28),('Careers Page','careers',36),('LinkedIn','linkedin',30),('Company Size','size',18),
     ('Cold Outreach Practical?','cold',44),('Sources','sources',66)]
ws2.append([c[0] for c in d])
for i, r in enumerate(K, 1):
    ws2.append([i] + [r.get(c[1], 'Not found') for c in d[1:]])
style(ws2, 1, len(d), [c[2] for c in d]); ws2.freeze_panes = 'C2'; ws2.auto_filter.ref = ws2.dimensions

# Sheet 3: emailable shortlist
ws3 = wb.create_sheet('Emailable Now (21)')
ws3['A1'] = ('The 21 companies with a publicly verified email. The 8 with a real HR/recruitment mailbox are first - '
             'those are the ones worth writing to first. No address here was guessed or built from a naming pattern.')
ws3.merge_cells('A1:H1'); ws3['A1'].font = Font(bold=True, size=11)
ws3['A1'].alignment = Alignment(wrap_text=True, vertical='center'); ws3.row_dimensions[1].height = 32
e = [r for r in K if r['_contactable'] == 'YES']
ws3.append(['No.','Company','Station','HR / Recruitment Email','General Email','Phone','Mailbox type','Score /6'])
for i, r in enumerate(e, 1):
    ws3.append([i, r['company'], r['station'], r.get('hr_email'), r.get('gen_email'), r.get('phone'),
                'HR / recruitment' if r['_f']['1 HR/recruitment email']=='YES' else 'General / business', r['_ready']])
style(ws3, 2, 8, [5,34,18,32,32,20,20,9]); ws3.freeze_panes = 'A3'

# Sheet 4: removals
ws4 = wb.create_sheet('Removed (25) + Why')
ws4['A1'] = 'The 25 entries cut from the original 100, with the criterion each one failed.'
ws4.merge_cells('A1:F1'); ws4['A1'].font = Font(bold=True, size=11)
ws4.row_dimensions[1].height = 22
ws4.append(['No.','Company','Criterion failed','Reason','Area as listed','Station'])
for i, r in enumerate(sorted(R, key=lambda x: (x['_criterion'], x['company'])), 1):
    ws4.append([i, r['company'], r['_criterion'], r['_why_removed'], r['area'], r['station']])
style(ws4, 2, 6, [5,40,22,74,40,16]); ws4.freeze_panes = 'A3'

wb.save('Navi_Mumbai_VERIFIED_75_Final.xlsx')
print('saved. kept', len(K), 'emailable', len(e), 'removed', len(R))
