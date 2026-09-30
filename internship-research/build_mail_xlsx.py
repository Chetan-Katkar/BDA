import sys,json; sys.path.insert(0,'.')
from new_mumbai import N
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter
DROP={"IBM India","Capgemini (Mindspace Malad)","Accenture (Mindspace Malad)","L&T Infotech (LTIMindtree Mumbai)"}
NEW=[n for n in N if n[0] not in DROP]
OLD=json.load(open('companies.json'))
KEPT={r['company'] for r in json.load(open('verified_final.json'))}
NF=lambda v:(not v) or str(v).strip().lower().startswith('not found')
def dom(s):
    import re
    d=re.findall(r'https?://([^/\s;]+)',str(s))
    return ', '.join(dict.fromkeys(d))[:70]
H=PatternFill('solid',fgColor='1F3864'); HF=Font(bold=True,color='FFFFFF',size=9)
wb=Workbook()
rm=wb.active; rm.title='READ ME'
for i,t in enumerate([
 "Mumbai + Navi Mumbai internship target list - 289 companies",
 "100 Navi Mumbai (individually researched) + 189 wider Mumbai (discovery level).",
 "",
 "THE TWO HALVES ARE NOT THE SAME QUALITY. Read this.",
 "Navi Mumbai 100: address, contacts, careers page and internship evidence chased per company.",
 "  75 passed a strict audit; 25 are marked FAILED AUDIT. 21 have a publicly verified email.",
 "  Use these for real outreach - see the Emailable Now sheet.",
 "Wider Mumbai 189: real companies from curated lists, review directories and job listings, with",
 "  sector and source. NO verified address, email, phone or internship status. Open the company",
 "  site before you contact any of them. MED = locality/reviewed-directory found. LOW = name only.",
 "",
 "WHAT I COULD NOT VERIFY, AND WHY",
 "Page fetching is blocked in the research environment, so Glassdoor/AmbitionBox rating pages could",
 "not be opened. There is NO invented rating or stipend anywhere in this file. Stipends appear only",
 "where a real listing stated one, e.g. Silicon Interfaces Rs 15,000-25,000/mo, Host360 up to Rs 12,000/mo.",
 "'Easy to crack' cannot be evidenced; company size plus a live intern posting is the honest proxy.",
 "Target was 250 new companies; I reached 189 distinct ones without padding the list.",
 "",
 "Easy + high stipend + great reviews rarely coincide. Dream11, BrowserStack, Google and the bank",
 "GCCs pay best and are hardest to enter. The small Navi Mumbai firms with a published hr@ mailbox",
 "are likeliest to reply. Apply to big names via portals; cold-email the small ones.",
 "",
 "Full-detail workbook (distances, roles, evidence, full source URLs, the 25 audit failures with",
 "reasons) is in the repo: github.com/Chetan-Katkar/BDA, branch claude/eloquent-babbage-rsuqq4,",
 "folder internship-research/",
],1):
    rm.cell(i,1,t).font=Font(size=11 if i in (1,4,12) else 10, bold=i in (1,4,12))
rm.column_dimensions['A'].width=112

ws=wb.create_sheet('All Companies')
cols=['No.','Company','List','Area / Locality','Zone / Station','Sector','Status','HR Email','General Email','Phone','Careers / Website','Conf','Source']
ws.append(cols); n=0
for r in OLD:
    n+=1
    ws.append([n,r['company'],'NM-100',r['area'],r['station'],r['domain'],
               'PASSED' if r['company'] in KEPT else 'FAILED AUDIT',
               r.get('hr_email'),r.get('gen_email'),r.get('phone'),
               r.get('careers') if not NF(r.get('careers')) else r.get('website'),'HIGH',dom(r.get('sources'))])
for nm,area,zone,sector,conf,src in NEW:
    n+=1
    ws.append([n,nm,'Mumbai',area,zone,sector,'DISCOVERY','-','-','-','-',conf,src])
for ci,w in enumerate([5,32,9,32,20,34,14,24,24,16,30,7,30],1):
    ws.column_dimensions[get_column_letter(ci)].width=w
    c=ws.cell(1,ci); c.fill=H; c.font=HF; c.alignment=Alignment(wrap_text=True,horizontal='center')
for row in ws.iter_rows(min_row=2,max_row=ws.max_row,max_col=len(cols)):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True); c.font=Font(size=9)
ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions

w4=wb.create_sheet('Emailable Now')
em=[r for r in OLD if r['company'] in KEPT and (not NF(r.get('hr_email')) or not NF(r.get('gen_email')))]
em.sort(key=lambda r: NF(r.get('hr_email')))
w4.append(['No.','Company','Station','HR Email','General Email','Phone','Type'])
for i,r in enumerate(em,1):
    w4.append([i,r['company'],r['station'],r.get('hr_email'),r.get('gen_email'),r.get('phone'),
               'HR' if not NF(r.get('hr_email')) else 'General'])
for ci,w in enumerate([5,34,20,32,32,20,10],1):
    w4.column_dimensions[get_column_letter(ci)].width=w
    c=w4.cell(1,ci); c.fill=H; c.font=HF; c.alignment=Alignment(wrap_text=True,horizontal='center')
for row in w4.iter_rows(min_row=2,max_row=w4.max_row,max_col=7):
    for c in row: c.alignment=Alignment(vertical='top',wrap_text=True); c.font=Font(size=9)
wb.save('Mumbai_Internship_List_289.xlsx')
import os,base64
sz=os.path.getsize('Mumbai_Internship_List_289.xlsx')
b=base64.b64encode(open('Mumbai_Internship_List_289.xlsx','rb').read()).decode()
open('mail_att.b64','w').write(b)
print('xlsx bytes:',sz,'| base64 chars:',len(b))
