import json, csv, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

D = json.load(open('companies.json'))
NF = lambda v: (not v) or str(v).strip().lower().startswith('not found')

STACK = re.compile(r'java|spring|python|sql|backend|rest|api|ai|machine learning|ml|cloud|devops|data engineer|full.?stack|mern|typescript|\.net', re.I)

def score(r):
    pts, why = 0, []
    op = str(r.get('opening',''))
    if op.strip().upper().startswith('YES'):
        pts += 3; why.append('current opening confirmed')
    if not NF(r.get('evidence')):
        pts += 2; why.append('documented intern/fresher hiring')
    if not NF(r.get('hr_email')):
        pts += 3; why.append('public HR/recruitment email')
    if not NF(r.get('gen_email')):
        pts += 1; why.append('public general email')
    if not NF(r.get('phone')):
        pts += 1; why.append('public phone')
    if not NF(r.get('careers')):
        pts += 1; why.append('active careers page')
    sz = str(r.get('size','')) + ' ' + str(r.get('type',''))
    if re.search(r'small|startup|very small|11-50', sz, re.I):
        pts += 2; why.append('small/startup - more approachable')
    if STACK.search(str(r.get('domain','')) + ' ' + str(r.get('roles',''))):
        pts += 1; why.append('stack matches your skills')
    tier = 'A-HIGH' if pts >= 9 else ('B-MEDIUM' if pts >= 6 else ('C-LOW' if pts >= 3 else 'D-THIN DATA'))
    return pts, tier, '; '.join(why) if why else 'location verified only'

for r in D:
    p, t, w = score(r)
    r['_score'], r['_tier'], r['_why'] = p, t, w

D.sort(key=lambda r: (-r['_score'], r['company']))

COLS = [
 ('No.', None, 5),
 ('Company', 'company', 34),
 ('Area / Office Address', 'area', 46),
 ('Nearest Station', 'station', 16),
 ('Approx. Distance', 'dist', 22),
 ('Company Type', 'type', 26),
 ('Technology / Domain', 'domain', 40),
 ('Relevant Internship Roles', 'roles', 40),
 ('Current Opening?', 'opening', 30),
 ('Internship Evidence', 'evidence', 46),
 ('HR / Recruitment Email', 'hr_email', 28),
 ('General Email', 'gen_email', 28),
 ('Phone', 'phone', 22),
 ('Website', 'website', 30),
 ('Careers Page', 'careers', 38),
 ('LinkedIn', 'linkedin', 30),
 ('Company Size', 'size', 20),
 ('Cold Outreach Practical?', 'cold', 46),
 ('Internship Outreach Priority', '_tier', 14),
 ('Priority Score', '_score', 10),
 ('Priority Basis (factual indicators)', '_why', 44),
 ('Source / Evidence Links', 'sources', 70),
]

HDR = PatternFill('solid', fgColor='1F3864')
HF  = Font(bold=True, color='FFFFFF', size=10)
THIN = Border(*[Side(style='thin', color='D0D0D0')]*4)
TIERC = {'A-HIGH':'C6EFCE','B-MEDIUM':'FFF2CC','C-LOW':'FCE4D6','D-THIN DATA':'EDEDED'}

wb = Workbook()

# --- Sheet 1: master table
ws = wb.active; ws.title = '100 Companies'
ws.append([c[0] for c in COLS])
for i, r in enumerate(D, 1):
    ws.append([i] + [r.get(c[1], 'Not found') if c[1] else i for c in COLS[1:]])
for ci, (name, key, w) in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
    cell = ws.cell(1, ci); cell.fill = HDR; cell.font = HF
    cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center')
ws.row_dimensions[1].height = 34
tier_col = [c[1] for c in COLS].index('_tier') + 1
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(COLS)):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True); c.border = THIN; c.font = Font(size=9)
    tv = row[tier_col-1].value
    if tv in TIERC:
        row[tier_col-1].fill = PatternFill('solid', fgColor=TIERC[tv])
        row[tier_col-1].font = Font(size=9, bold=True)
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions

# --- Section sheets A / B / C
A = [r for r in D if str(r.get('opening','')).strip().upper().startswith('YES')]
B = [r for r in D if r not in A and not NF(r.get('evidence'))]
C = [r for r in D if r not in A and r not in B]

def section(title, rows, cols, note):
    s = wb.create_sheet(title)
    s['A1'] = note; s['A1'].font = Font(bold=True, size=11)
    s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    s['A1'].alignment = Alignment(wrap_text=True, vertical='center'); s.row_dimensions[1].height = 30
    s.append([c[0] for c in cols])
    for i, r in enumerate(rows, 1):
        s.append([i] + [r.get(c[1], 'Not found') for c in cols[1:]])
    for ci, (n, k, w) in enumerate(cols, 1):
        s.column_dimensions[get_column_letter(ci)].width = w
        hc = s.cell(2, ci); hc.fill = HDR; hc.font = HF
        hc.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    s.row_dimensions[2].height = 30
    for row in s.iter_rows(min_row=3, max_row=s.max_row, max_col=len(cols)):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True); c.border = THIN; c.font = Font(size=9)
    s.freeze_panes = 'A3'
    return s

SEC = [('No.',None,5),('Company','company',32),('Area','area',42),('Nearest Station','station',16),
       ('Current Opening?','opening',30),('Internship Evidence','evidence',46),
       ('HR / Recruitment Email','hr_email',28),('General Email','gen_email',26),('Phone','phone',20),
       ('Careers Page','careers',36),('Relevant Roles','roles',38),('Priority','_tier',14),
       ('Why relevant (factual)','_why',42),('Sources','sources',60)]

section('A - Current Openings', A,  SEC, 'SECTION A - Companies with a CURRENT internship / entry-level opening found in a public posting at research time (30 Sep 2026). Job postings expire: re-check each link before applying.')
section('B - Intern-Fresher Evidence', B, SEC, 'SECTION B - No specific current opening confirmed, but there is documented public evidence that the company hires interns and/or freshers (named internship programme, internship-platform employer profile, walk-in drives, or an active careers pipeline).')
section('C - Worth Contacting', C, SEC, 'SECTION C - No internship evidence found in public sources. Listed because the office location in your target corridor is verified and the technology domain is relevant. The "Why relevant (factual)" column states the factual basis only - it is not a prediction that they will hire you.')

# --- Verified-email outreach shortlist
em = [r for r in D if not NF(r.get('hr_email')) or not NF(r.get('gen_email'))]
em.sort(key=lambda r: (NF(r.get('hr_email')), -r['_score']))
s = wb.create_sheet('Verified Email Shortlist')
s['A1'] = 'Companies with a PUBLICLY VERIFIED email address. HR/recruitment mailboxes are listed first - start your outreach with those. No address on this sheet was guessed or built from a naming pattern.'
s.merge_cells('A1:H1'); s['A1'].font = Font(bold=True, size=11)
s['A1'].alignment = Alignment(wrap_text=True, vertical='center'); s.row_dimensions[1].height = 32
ecols = [('No.',None,5),('Company','company',34),('Nearest Station','station',16),
         ('HR / Recruitment Email','hr_email',30),('General Email','gen_email',30),
         ('Phone','phone',22),('Email type','_etype',22),('Source','sources',70)]
for r in em:
    r['_etype'] = 'HR / recruitment mailbox' if not NF(r.get('hr_email')) else 'General / business mailbox'
s.append([c[0] for c in ecols])
for i, r in enumerate(em, 1):
    s.append([i] + [r.get(c[1], 'Not found') for c in ecols[1:]])
for ci, (n, k, w) in enumerate(ecols, 1):
    s.column_dimensions[get_column_letter(ci)].width = w
    hc = s.cell(2, ci); hc.fill = HDR; hc.font = HF
    hc.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
for row in s.iter_rows(min_row=3, max_row=s.max_row, max_col=len(ecols)):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True); c.border = THIN; c.font = Font(size=9)
s.freeze_panes = 'A3'

# --- Outreach tracker
t = wb.create_sheet('OUTREACH TRACKER')
t['A1'] = 'OUTREACH TRACKER - pre-filled with every company that has a publicly verified email, ordered by Internship Outreach Priority. Nothing has been sent. Fill in the Date Contacted column only after you send.'
t.merge_cells('A1:I1'); t['A1'].font = Font(bold=True, size=11)
t['A1'].alignment = Alignment(wrap_text=True, vertical='center'); t.row_dimensions[1].height = 32
tcols = [('Company',34),('Contact Person',22),('Email',32),('Date Contacted',16),('Follow-up Date',16),
         ('Response',18),('Interview',14),('Status',18),('Notes',46)]
t.append([c[0] for c in tcols])
for r in em:
    mail = r['hr_email'] if not NF(r.get('hr_email')) else r['gen_email']
    t.append([r['company'], '', str(mail).split(' ;')[0].strip(), '', '', '', '', 'Not contacted',
              f"{r['_tier']} | {r['station']} | {r['_why']}"])
for ci, (n, w) in enumerate(tcols, 1):
    t.column_dimensions[get_column_letter(ci)].width = w
    hc = t.cell(2, ci); hc.fill = HDR; hc.font = HF
    hc.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
for row in t.iter_rows(min_row=3, max_row=t.max_row, max_col=len(tcols)):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True); c.border = THIN; c.font = Font(size=9)
t.freeze_panes = 'A3'
for extra in range(t.max_row+1, t.max_row+26):
    for ci in range(1, len(tcols)+1):
        t.cell(extra, ci).border = THIN

# --- Read me
rm = wb.create_sheet('READ ME', 0)
lines = [
 ('Navi Mumbai internship research - 100 companies', 14, True),
 ('Compiled 30 September 2026. Target: a 2-6 month software/technology internship reachable from the Harbour and Trans-Harbour lines.', 10, False),
 ('', 10, False),
 ('SHEETS', 12, True),
 ('100 Companies - the full table, all 22 columns, sorted by Internship Outreach Priority (filters are on row 1).', 10, False),
 ('A - Current Openings - a public internship/entry-level posting existed at research time.', 10, False),
 ('B - Intern-Fresher Evidence - documented intern/fresher hiring, no specific current posting confirmed.', 10, False),
 ('C - Worth Contacting - verified office location and relevant domain, no internship evidence found.', 10, False),
 ('Verified Email Shortlist - every publicly verified email address, HR mailboxes first. Start here.', 10, False),
 ('OUTREACH TRACKER - pre-filled with the emailable companies. Nothing has been sent yet.', 10, False),
 ('', 10, False),
 ('HOW "Internship Outreach Priority" IS CALCULATED', 12, True),
 ('Factual indicators only - no opinion about company quality is involved:', 10, False),
 ('  +3  a current internship/entry-level opening was found in a public posting', 10, False),
 ('  +3  a publicly verified HR / recruitment email exists', 10, False),
 ('  +2  documented evidence of hiring interns or freshers', 10, False),
 ('  +2  small company or startup (smaller firms answer direct approaches more often)', 10, False),
 ('  +1  each: public general email, public phone, active careers page, domain matching your stack', 10, False),
 ('  Tiers: A-HIGH >=9, B-MEDIUM 6-8, C-LOW 3-5, D-THIN DATA <3', 10, False),
 ('', 10, False),
 ('VERIFICATION RULES APPLIED', 12, True),
 ('No email address was invented or reconstructed from a naming pattern. Where only a pattern was available', 10, False),
 ('(for example from LeadIQ, RocketReach, SalezShark or SignalHire), the field says "Not found" and the record says so.', 10, False),
 ('Every row carries at least one source link. "Not found" means genuinely not publicly verifiable, not "not checked".', 10, False),
 ('Addresses are marked where they come from a building tenant listing rather than the company\'s own site.', 10, False),
 ('Registered offices are distinguished from operating offices where the two differ.', 10, False),
 ('', 10, False),
 ('KNOWN LIMITS - PLEASE READ', 12, True),
 ('1. Distances are map estimates from the station to the building, not walking-route measurements.', 10, False),
 ('2. Job postings expire. Every "Current Opening?" entry must be re-checked on the linked page before you apply.', 10, False),
 ('3. Details were gathered through web search result summaries; the research environment blocked direct page', 10, False),
 ('   fetching, so first-party pages could not be opened and read line by line. Confirm each email on the', 10, False),
 ('   company contact page before sending, especially the general/business mailboxes.', 10, False),
 ('4. Excluded on purpose: Nipralo Technologies (office is Ghatkopar East, outside the corridor),', 10, False),
 ('   Skypal System (listed as closed down), Infosys Ltd (no verifiable Ghansoli office),', 10, False),
 ('   and pure training institutes such as Netweaver Technovations and PNK IT Solutions.', 10, False),
 ('5. Some entries are captive/GCC offices of banks and shipping firms. They do employ software teams and their', 10, False),
 ('   locations are verified, but they are clearly labelled so you can judge the fit yourself.', 10, False),
]
for i, (txt, sz, bold) in enumerate(lines, 1):
    c = rm.cell(i, 1, txt); c.font = Font(size=sz, bold=bold)
    c.alignment = Alignment(wrap_text=False, vertical='center')
rm.column_dimensions['A'].width = 125

wb.save('Navi_Mumbai_100_Companies_Internship_Research.xlsx')

# --- CSVs
with open('companies_100.csv','w',newline='',encoding='utf-8') as f:
    w = csv.writer(f); w.writerow([c[0] for c in COLS])
    for i, r in enumerate(D,1):
        w.writerow([i] + [r.get(c[1],'Not found') for c in COLS[1:]])
with open('outreach_tracker.csv','w',newline='',encoding='utf-8') as f:
    w = csv.writer(f); w.writerow([c[0] for c in tcols])
    for r in em:
        mail = r['hr_email'] if not NF(r.get('hr_email')) else r['gen_email']
        w.writerow([r['company'],'',str(mail).split(' ;')[0].strip(),'','','','','Not contacted',
                    f"{r['_tier']} | {r['station']} | {r['_why']}"])

print('companies:', len(D))
print('A current openings:', len(A), '| B evidence:', len(B), '| C contact-worthy:', len(C))
print('with verified email:', len(em), '| of which HR/recruitment mailbox:', sum(1 for r in em if not NF(r.get('hr_email'))))
from collections import Counter
print('tiers:', dict(Counter(r['_tier'] for r in D)))
print('stations:', dict(Counter(r['station'].split(' /')[0].split(' (')[0].strip() for r in D)))
